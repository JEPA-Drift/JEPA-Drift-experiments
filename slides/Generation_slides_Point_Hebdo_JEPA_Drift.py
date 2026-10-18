"""
Build PPTX for Point Hebdo C.S. du 12/10/2026 — JEPA-Drift : l'idée du nouvel article
Same visual style as 2026_04_20_Point_Hebdo.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from copy import deepcopy
import os

# === COLOR PALETTE (from reference deck) ===
TEAL_DARK   = RGBColor(0x00, 0x74, 0x8C)  # primary brand
TEAL        = RGBColor(0x1B, 0x7A, 0x8A)  # header bar
BLUE_DEEP   = RGBColor(0x3B, 0x6F, 0xA0)
BLUE        = RGBColor(0x29, 0x80, 0xB9)
ORANGE      = RGBColor(0xE0, 0x80, 0x00)
ORANGE_LT   = RGBColor(0xE8, 0x83, 0x3A)
RED         = RGBColor(0xC0, 0x39, 0x2B)
RED_LT      = RGBColor(0xE7, 0x4C, 0x3C)
PURPLE      = RGBColor(0x8E, 0x44, 0xAD)
GREEN       = RGBColor(0x27, 0xAE, 0x60)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BG_BLUE_LT  = RGBColor(0xEB, 0xF2, 0xFA)
BG_CREAM    = RGBColor(0xFF, 0xF8, 0xF0)
BG_PINK     = RGBColor(0xFF, 0xF0, 0xF0)
BG_GREEN_LT = RGBColor(0xE8, 0xF6, 0xEF)
TXT_DARK    = RGBColor(0x2C, 0x3E, 0x50)
TXT_BODY    = RGBColor(0x5D, 0x6D, 0x7E)
TXT_FAINT   = RGBColor(0x8A, 0x9B, 0xB0)
TXT_DARKER  = RGBColor(0x2D, 0x37, 0x48)

# === SLIDE GEOMETRY ===
SLIDE_W = 10.00  # inches
SLIDE_H = 5.62   # inches
TOTAL_SLIDES = 12  # will be set after construction
FIG_DIR = "figures"

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)

# Use blank layout
BLANK = prs.slide_layouts[6]

def add_rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s

def add_text(slide, x, y, w, h, text, font_name="Calibri", size=11, bold=False,
             color=TXT_DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb

def add_multi_text(slide, x, y, w, h, runs, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """runs = list of (text, size, bold, color, italic_optional)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for i, r in enumerate(runs):
        text = r[0]
        size = r[1]
        bold = r[2]
        color = r[3]
        italic = r[4] if len(r) > 4 else False
        if i == 0:
            run = p.add_run()
        else:
            run = p.add_run()
        run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return tb

def add_paragraphs(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP):
    """paragraphs = list of list of (text, size, bold, color, italic_optional)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for pi, para in enumerate(paragraphs):
        if pi == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        for r in para:
            text = r[0]
            size = r[1]
            bold = r[2]
            color = r[3]
            italic = r[4] if len(r) > 4 else False
            run = p.add_run()
            run.text = text
            run.font.name = "Calibri"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
    return tb

def header_bar(slide, title, page_num=None, total=20):
    idx = len(prs.slides)
    add_rect(slide, 0, 0, 10.00, 0.58, TEAL)
    add_text(slide, 0.40, 0.05, 8.50, 0.50, title, size=17, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 8.80, 5.20, 1.00, 0.35, f"{idx}", size=9, color=TXT_BODY,
             align=PP_ALIGN.RIGHT)

def page_num_only(slide, page_num=None, total=20):
    idx = len(prs.slides)
    add_text(slide, 8.80, 5.20, 1.00, 0.35, f"{idx}", size=9, color=TXT_BODY,
             align=PP_ALIGN.RIGHT)

def badge(slide, text, x=8.90, y=0.07, w=0.90, h=0.27, color=RED):
    add_rect(slide, x, y, w, h, color)
    add_text(slide, x, y, w, h, text, size=8, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def vertical_bar(slide, x, y, color=TEAL, height=0.70):
    add_rect(slide, x, y, 0.06, height, color)

def lettered_circle(slide, x, y, letter, color=TEAL, size=0.55, fontsize=18):
    add_rect(slide, x, y, size, size, color, shape=MSO_SHAPE.OVAL)
    add_text(slide, x, y, size, size, letter, size=fontsize, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def footer_refs(slide, refs):
    add_text(slide, 0.50, 5.36, 8.00, 0.26, refs, size=8, color=TXT_FAINT)

# =========================================================================
# SLIDE 1 — TITRE
# =========================================================================
s = prs.slides.add_slide(BLANK)
# Top teal block
add_rect(s, 0.60, 0.50, 8.80, 2.20, TEAL_DARK)
add_paragraphs(s, 0.75, 0.65, 8.50, 2.00, [
    [("JEPA-Drift", 28, True, WHITE)],
    [("Escaping the Sharpe Ceiling", 28, True, WHITE)],
    [("of Sequential Drift Detection", 26, True, WHITE)],
    [("", 8, False, WHITE)],
    [("through Latent Information Concentration", 16, False, WHITE, True)],
])
# Orange accent line
add_rect(s, 0.60, 2.80, 8.80, 0.06, ORANGE)
# Lower blue strip
add_rect(s, 0.00, 3.00, 10.00, 1.80, BLUE_DEEP)
# Orange separator
add_rect(s, 0.60, 3.10, 8.80, 0.46, ORANGE)
add_text(s, 0.70, 3.14, 8.60, 0.38, "Point hebdo du lundi 12/10/2026 — Article C : l'idée du nouvel article",
         size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, 0.70, 3.62, 8.60, 0.38, "R. Minato — suite de « The Whitening Advantage » (Article B, soumis à KDD 2027)",
         size=12, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
idx = len(prs.slides)
add_text(s, 9.00, 5.25, 0.80, 0.30, f"{idx}", size=9, color=WHITE,
         align=PP_ALIGN.RIGHT)

# =========================================================================
# SLIDE 2 — Le point de départ : le plancher de Sharpe
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Le point de départ : le plancher de Sharpe", 2)
badge(s, "CONSTAT", color=RED, w=0.95)

# Deux colonnes
# Colonne gauche (constat)
add_rect(s, 0.30, 0.75, 4.45, 0.35, BG_PINK)
add_text(s, 0.30, 0.75, 4.45, 0.35, "  Détecter une dérive, c'est accumuler des preuves",
         size=10, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.35, 1.20, 4.40, 1.50, [
    [("• Chaque jour apporte un peu de preuve qu'une dérive a commencé. On la mesure par la divergence de Kullback-Leibler (KL) entre la loi des rendements avant et après la rupture.", 9, False, TXT_BODY)],
    [("• Un détecteur conclut quand la preuve cumulée franchit un seuil. Le délai moyen de détection (ADD) vaut donc, au mieux, seuil ÷ preuve quotidienne.", 9, False, TXT_BODY)],
    [("• Le seuil traduit la tolérance aux fausses alarmes : ln(1/α₀), où α₀ = 1/ARL₀ et ARL₀ est le nombre moyen de jours entre deux fausses alarmes.", 9, False, TXT_BODY)],
    [("• Pour une petite dérive μ de la moyenne journalière, à volatilité σ, la preuve quotidienne vaut μ²/(2σ²) : la moitié du carré du Sharpe journalier μ/σ.", 9, False, TXT_BODY)],
])

# Colonne droite (coût en jours)
add_rect(s, 5.25, 0.75, 4.45, 0.35, BG_GREEN_LT)
add_text(s, 5.25, 0.75, 4.45, 0.35, "  Ce que cela coûte en jours",
         size=10, bold=True, color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 5.30, 1.20, 4.40, 1.80, [
    [("• En annualisant (252 jours ; SR annuel = √252 · μ/σ), la preuve quotidienne vaut SR²/504.", 9, False, TXT_BODY)],
    [("• SR = 1, une fausse alarme par an (ARL₀ = 252) : au moins 2 800 jours, soit 11 ans. SR = 0,5 : 44 ans.", 9, False, TXT_BODY)],
    [("• Avec une fausse alarme par mois (ARL₀ = 20), SR = 1 coûte encore 1 500 jours, soit 6 ans.", 9, False, TXT_BODY)],
    [("• Sur 66 phases haussières et baissières datées (4 ETF, 2000–2025) : 80 % durent moins que leur plancher même avec une fausse alarme par mois ; 97 % avec une fausse alarme par an.", 9, False, TXT_BODY)],
    [("• Sur une série seule, aucun algorithme n'y échappe : il manque de l'information, pas de la méthode.", 9, False, TXT_BODY)],
])

# Encadré formule
add_rect(s, 0.30, 3.30, 9.40, 1.40, BG_BLUE_LT)
add_text(s, 0.30, 3.30, 9.40, 0.30, "  Le plancher de Sharpe (Article B)", size=10, bold=True,
         color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.40, 3.65, 9.20, 1.00, [
    [("ADD_min ≥ 504 · ln(1/α₀) / SR²  (jours de bourse)", 9, False, TXT_DARKER, True)],
    [("", 4, False, TXT_DARKER)],
    [("• Diviser le Sharpe par deux multiplie le délai par quatre.", 9, False, TXT_BODY)],
])

# VISUEL À DESSINER À LA MAIN : ADD_min en années (échelle log) en fonction du SR annuel de 0,25 à 3, deux courbes (ARL₀ = 20 et 252), bande grise des durées typiques de phases (quelques mois à quelques années).

footer_refs(s, "Lorden (1971) · Lai (1998) · Tartakovsky et al. (2014) · Pagan & Sossounov (2003)")

# =========================================================================
# SLIDE 3 — Additionner des actifs corrélés ne suffit pas
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Additionner des actifs corrélés ne suffit pas", 3)
badge(s, "CONSTAT", color=RED, w=0.95)

# Colonne gauche (pooling)
add_rect(s, 0.30, 0.75, 4.45, 0.35, BG_PINK)
add_text(s, 0.30, 0.75, 4.45, 0.35, "  Le pooling : additionner les signaux de K actifs",
         size=10, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.35, 1.20, 4.40, 1.50, [
    [("• Idée naturelle : surveiller K actions à la fois et additionner leurs signaux.", 9, False, TXT_BODY)],
    [("• Avec K séries indépendantes, la preuve quotidienne serait multipliée par K.", 9, False, TXT_BODY)],
    [("• Sur 97 actions américaines (2005–2025), les signes des rendements sont corrélés : ρ̂ ≈ 0,26. Le panel vaut K_eff ≈ 1/ρ̂ ≈ 3,8 actions indépendantes, quel que soit K.", 9, False, TXT_BODY)],
    [("• Mesuré : le budget de détection ne baisse que d'un facteur 2, la limite √K_eff, et plafonne dès K ≈ 20.", 9, False, TXT_BODY)],
    [("• Supposer l'indépendance fait grimper les fausses alarmes vers 100 % dès K = 40 ; un bootstrap sur fenêtres réelles tient les 5 %.", 9, False, TXT_BODY)],
    [("• Le pooling ne regarde qu'une direction, la moyenne du panel, c'est-à-dire le marché : le mur s'élargit, il ne tombe pas.", 9, False, TXT_BODY)],
])

# Colonne droite : figure (Article B, Fig. 30)
add_rect(s, 5.25, 0.75, 4.45, 0.35, BG_GREEN_LT)
add_text(s, 5.25, 0.75, 4.45, 0.35, "  Mesure réelle (Article B, Fig. 30)",
         size=10, bold=True, color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_picture(os.path.join(FIG_DIR, "Fig30_RealCrossSectional_Escape_UPDATED.png"), Inches(5.25), Inches(1.15), width=Inches(4.45))

# Encadré formule
add_rect(s, 0.30, 3.30, 9.40, 1.40, BG_BLUE_LT)
add_text(s, 0.30, 3.30, 9.40, 0.30, "  Pourquoi ça plafonne", size=10, bold=True,
         color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.40, 3.65, 9.20, 1.00, [
    [("Var(moyenne de K actifs) = σ² · (ρ + (1 − ρ)/K)  ;  K_eff = K / (1 + (K − 1)·ρ) → 1/ρ ≈ 3,8", 9, False, TXT_DARKER, True)],
    [("", 4, False, TXT_DARKER)],
    [("• ρ : corrélation moyenne entre deux flux surveillés, ici les signes. La part commune ne se diversifie jamais.", 9, False, TXT_BODY)],
])

footer_refs(s, "Article B, Fig. 30 · 97 actions US survivantes, 2005–2025")

# =========================================================================
# SLIDE 4 — L'idée : chercher la dérive là où le marché ne fait pas de bruit
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "L'idée : chercher la dérive là où le marché ne fait pas de bruit", 4)
badge(s, "IDÉE", color=GREEN, w=0.90)

# Colonne gauche (mode de marché)
add_rect(s, 0.30, 0.75, 4.45, 0.35, BG_PINK)
add_text(s, 0.30, 0.75, 4.45, 0.35, "  Ce qui noie la dérive : le mode de marché",
         size=10, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.35, 1.20, 4.40, 1.50, [
    [("• On surveille N actifs à la fois : chaque jour, un vecteur de N rendements.", 9, False, TXT_BODY)],
    [("• Sa covariance Σ se découpe en directions orthogonales v₁, …, v_N (composantes principales), de variances λ₁ ≥ … ≥ λ_N.", 9, False, TXT_BODY)],
    [("• La première, le marché, porte 36 à 40 % de la variance de nos panels S&P 500, davantage en crise.", 9, False, TXT_BODY)],
    [("• Une dérive logée dans cette direction se heurte à tout ce bruit : c'est le sort du pooling.", 9, False, TXT_BODY)],
])

# Colonne droite (information)
add_rect(s, 5.25, 0.75, 4.45, 0.35, BG_GREEN_LT)
add_text(s, 5.25, 0.75, 4.45, 0.35, "  Où se loge l'information",
         size=10, bold=True, color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 5.30, 1.20, 4.40, 1.80, [
    [("• Une dérive Δμ, déplacement du vecteur des moyennes, apporte dans chaque direction une preuve égale à (part de Δμ dans cette direction)² ÷ variance de la direction.", 9, False, TXT_BODY)],
    [("• À taille égale, une dérive dans une direction calme pèse bien plus qu'une dérive alignée sur le marché.", 9, False, TXT_BODY)],
    [("• Projeter ne crée aucune information (inégalité de traitement des données) : on retire le bruit qui la masquait.", 9, False, TXT_BODY)],
    [("• Si la dérive touche beaucoup de titres dans des directions calmes, les preuves s'additionnent : le gain peut croître avec le nombre de titres au lieu de plafonner.", 9, False, TXT_BODY)],
])

# Encadré formule
add_rect(s, 0.30, 3.30, 9.40, 1.40, BG_BLUE_LT)
add_text(s, 0.30, 3.30, 9.40, 0.30, "  La preuve quotidienne, direction par direction", size=10, bold=True,
         color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.40, 3.65, 9.20, 1.00, [
    [("KL = ½ · Σ_k (v_kᵀ Δμ)² / λ_k", 9, False, TXT_DARKER, True)],
    [("", 4, False, TXT_DARKER)],
    [("• Chaque direction pèse en 1/λ_k : à dérive égale, une direction dix fois plus calme pèse dix fois plus.", 9, False, TXT_BODY)],
])

# VISUEL À DESSINER À LA MAIN : nuage de rendements en ellipse allongée, grand axe étiqueté « marché » ; deux flèches de même longueur depuis le centre, l'une sur le grand axe (« noyée »), l'autre sur le petit axe (« visible »).

footer_refs(s, "Kullback (1959) · Cover & Thomas (2006) · Alouadi & Lehalle (2026)")

# =========================================================================
# SLIDE 5 — Deux façons de construire l'espace latent
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Deux façons de construire l'espace latent", 5)
badge(s, "MÉTHODE", color=BLUE, w=1.00)

# Two columns: référence linéaire / JEPA
col_w = 4.45
# Colonne gauche (PCA-orthogonal)
add_rect(s, 0.30, 0.75, col_w, 0.35, BG_PINK)
add_text(s, 0.30, 0.75, col_w, 0.35, "  Référence linéaire : PCA-orthogonal", size=10, bold=True,
         color=RED, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.40, 1.15, col_w - 0.10, 1.40, [
    [("• On calcule les composantes principales (PCA) du panel sur une fenêtre passée.", 9, False, TXT_BODY)],
    [("• On écarte la première, le marché, et on surveille les directions restantes.", 9, False, TXT_BODY)],
    [("• Linéaire, interprétable, rapide : c'est la référence à battre.", 9, False, TXT_BODY)],
    [("• Limite : directions fixes, choisies pour leur variance et non pour la dérive ; aveugle aux effets non linéaires (queues épaisses, volatilité commune).", 9, False, TXT_BODY)],
])

# Colonne droite (JEPA)
add_rect(s, 5.25, 0.75, col_w, 0.35, BG_GREEN_LT)
add_text(s, 5.25, 0.75, col_w, 0.35, "  Représentation apprise : JEPA", size=10, bold=True,
         color=GREEN, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 5.35, 1.15, col_w - 0.10, 1.40, [
    [("• JEPA (Joint-Embedding Predictive Architecture) : on masque une partie des actifs, et un réseau prédit leur représentation à partir des autres, sans reconstruire les données brutes.", 9, False, TXT_BODY)],
    [("• T-JEPA pour les données tabulaires ; CA-T-JEPA (cross-asset) pour la coupe d'actions, chaque actif devenant un jeton décrit par ses caractéristiques.", 9, False, TXT_BODY)],
    [("• Prédire dans l'espace latent favorise les directions prévisibles plutôt que les plus variables (démontré pour des réseaux linéaires profonds).", 9, False, TXT_BODY)],
    [("• Ce qu'il ne verra pas : une dérive propre à un seul titre, imprévisible depuis les autres.", 9, False, TXT_BODY)],
])

# Bandeau — comparaison à armes égales
add_rect(s, 0.30, 2.75, 9.40, 0.35, BG_BLUE_LT)
add_text(s, 0.30, 2.75, 9.40, 0.35, "  Une comparaison à armes égales",
         size=10, bold=True, color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)

calibs = [
    ("Même moniteur", "surveillance séquentielle",
     "les deux espaces alimentent les mêmes détecteurs, calibrés de la même façon."),
    ("Même dimension", "espace compact",
     "quelques dimensions seulement : chaque dimension surveillée ajoute son propre bruit."),
    ("Même bruit de fond", "données réelles",
     "mêmes panels, mêmes dérives injectées, mêmes dates de rupture."),
]
for i, (lam, sec, desc) in enumerate(calibs):
    y = 3.20 + i * 0.55
    vertical_bar(s, 0.40, y, color=BLUE, height=0.45)
    add_multi_text(s, 0.65, y, 9.10, 0.45,[
        (lam + "  ", 10, True, BLUE),
        ("(" + sec + ")  — ", 9, False, TXT_DARK),
        (desc, 9, False, TXT_BODY),
    ], anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, 0.30, 4.78, 9.40, 0.32, BG_CREAM)
add_text(s, 0.40, 4.78, 9.30, 0.32,
         "★ À dimension égale, la représentation apprise retient-elle plus d'information de dérive que la projection linéaire ?",
         size=10, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)

# VISUEL À DESSINER À LA MAIN : panel de rendements → (a) PCA, direction marché retirée | (b) encodeur JEPA → espace latent compact → même moniteur → alarme.

footer_refs(s, "Assran et al. (2023) · Thimonier et al. (2025) · Littwin et al. (2024)")

# =========================================================================
# SLIDE 6 — Ce qui existe, et ce qui manque
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Ce qui existe, et ce qui manque", 6)
badge(s, "ANALYSE", color=BLUE, w=1.00)

# Table header
add_rect(s, 0.30, 0.75, 9.40, 0.35, TEAL)
headers = [(0.30, 2.20, "Méthode"),
           (2.50, 1.40, "Représentation"),
           (3.90, 1.55, "Objectif"),
           (5.45, 1.55, "Détection séquentielle"),
           (7.00, 2.70, "Limite pour notre question")]
for x, w, text in headers:
    add_text(s, x, 0.75, w, 0.35, text, size=10, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Rows
data_rows = [
    ("PCA / Sparse PCA", "Linéaire", "Variance", "Non", "Facteurs statiques, non économiques", BG_BLUE_LT, BLUE),
    ("IPCA (2019)", "Linéaire, caract.", "Valorisation", "Non", "Lien linéaire caract. → expositions", BG_BLUE_LT, BLUE),
    ("GKX (2021)", "Non linéaire", "Valorisation", "Non", "Latent opaque, visé pour le pricing", BG_BLUE_LT, BLUE),
    ("CD-DFM (2026)", "Non lin., caract.", "Covariance future", "Non", "Visé pour le risque, pas la dérive", BG_BLUE_LT, BLUE),
    ("T-JEPA (2025)", "Non lin., JEPA", "Représentation", "Non", "Jamais testé en détection", BG_BLUE_LT, BLUE),
    ("Notre approche", "JEPA, compact", "Détection de dérive", "Oui, calibrée", "—", BG_GREEN_LT, GREEN),
]
for i, (conf, sub, reg, ov, profile, bg, color_text) in enumerate(data_rows):
    y = 1.10 + i * 0.32
    add_rect(s, 0.30, y, 9.40, 0.32, bg)
    add_text(s, 0.35, y, 2.20, 0.32, conf, size=9, bold=True, color=color_text,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 2.50, y, 1.40, 0.32, sub, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 3.90, y, 1.55, 0.32, reg, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 5.45, y, 1.55, 0.32, ov, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 7.00, y, 2.70, 0.32, profile, size=9, italic=True, color=TXT_BODY,
             anchor=MSO_ANCHOR.MIDDLE)

# Bloc de lecture
add_rect(s, 0.30, 3.10, 9.40, 0.32, BG_CREAM)
add_text(s, 0.30, 3.10, 9.40, 0.32, "  Lecture : ce qui est nouveau, et ce qui ne l'est pas",
         size=10, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.35, 3.50, 9.30, 1.65, [
    [("• Déjà connu : les changements de distribution ressortent davantage dans les composantes de faible variance d'une PCA (Kuncheva & Faithfull, 2014 ; Qahtan et al., 2015). Ce n'est pas notre revendication.", 9, False, TXT_BODY)],
    [("• Nouveau : une borne de délai dans l'espace latent, qui prolonge le plancher de Sharpe et chiffre ce qu'un espace latent peut gagner.", 9, False, TXT_BODY)],
    [("• Nouveau : une calibration exacte des moniteurs latents, valable quelle que soit la volatilité.", 9, False, TXT_BODY)],
    [("• Nouveau : une représentation apprise jugée contre ces références, à dimension et à fausses alarmes égales.", 9, False, TXT_BODY)],
    [("• ★ Au-delà de la finance : toute surveillance de capteurs corrélés (fraude, réseaux, industrie) bute sur le même mur.", 9, True, GREEN)],
])

footer_refs(s, "Kelly et al. (2019) · Gu et al. (2021) · Alouadi & Lehalle (2026) · Thimonier et al. (2025) · Kuncheva & Faithfull (2014) · Qahtan et al. (2015)")

# =========================================================================
# SLIDE 7 — Les grandes étapes
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Les grandes étapes", 7)
badge(s, "ÉTAPES", color=TEAL, w=0.95)

calibs = [
    ("Étape 1", "Théorie",
     "Borne de délai dans l'espace latent, concentration de l'information, taxe dimensionnelle."),
    ("Étape 2", "Audit des données",
     "Vérifier que l'information de dérive existe dans les panels financiers réels, et où elle se loge."),
    ("Étape 3", "Pile de surveillance",
     "Moniteurs séquentiels à calibration exacte, quelle que soit la volatilité."),
    ("Étape 4", "Référence linéaire",
     "Benchmark PCA-orthogonal sur bruit de fond réel."),
    ("Étape 5", "Représentations JEPA",
     "Pré-entraînement et contrôle de la qualité de l'espace latent."),
    ("Étape 6", "Confrontation",
     "JEPA contre PCA-orthogonal, sur données réelles et semi-synthétiques."),
    ("Étape 7", "Second domaine et rédaction",
     "Validation sur la fraude bancaire, puis rédaction."),
]
for i, (lam, sec, desc) in enumerate(calibs):
    y = 0.80 + i * 0.58
    vertical_bar(s, 0.40, y, color=BLUE, height=0.45)
    add_multi_text(s, 0.65, y, 9.10, 0.45,[
        (lam + "  ", 10, True, BLUE),
        ("(" + sec + ")  — ", 9, False, TXT_DARK),
        (desc, 9, False, TXT_BODY),
    ], anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, 0.30, 4.85, 9.40, 0.32, BG_CREAM)
add_text(s, 0.40, 4.85, 9.30, 0.32,
         "★ Sans moniteurs calibrés, aucun gain d'espace latent ne se mesure : la surveillance passe avant JEPA.",
         size=10, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)

# =========================================================================
# SLIDE 8 — Données et protocole de dérive
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Données et protocole de dérive", 8)
badge(s, "PROTOCOLE", color=BLUE, w=1.10)

# Table header
add_rect(s, 0.30, 0.75, 9.40, 0.35, TEAL)
headers = [(0.30, 2.20, "Jeu de données"),
           (2.50, 1.40, "Domaine"),
           (3.90, 1.55, "Taille"),
           (5.45, 1.55, "Dérives"),
           (7.00, 2.70, "Rôle")]
for x, w, text in headers:
    add_text(s, x, 0.75, w, 0.35, text, size=10, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Rows
data_rows = [
    ("Simulations", "Contrôle", "Illimitée", "Injectées", "Vérifier les bornes", BG_BLUE_LT, BLUE),
    ("S&P 500", "Actions US", "239–500 titres", "Réelles + injectées", "Domaine principal", BG_GREEN_LT, GREEN),
    ("EUROFIDAI", "Actions Europe", "Panel N × T", "Réelles + injectées", "Grands N, sans survivants", BG_GREEN_LT, GREEN),
    ("BAF (fraude)", "Comptes bancaires", "≈ 1 M demandes", "Mensuelles", "Second domaine", BG_BLUE_LT, BLUE),
    ("Électricité / IoT", "Énergie, capteurs", "—", "Réelles", "Option hors finance", BG_BLUE_LT, BLUE),
]
for i, (conf, sub, reg, ov, profile, bg, color_text) in enumerate(data_rows):
    y = 1.10 + i * 0.32
    add_rect(s, 0.30, y, 9.40, 0.32, bg)
    add_text(s, 0.35, y, 2.20, 0.32, conf, size=9, bold=True, color=color_text,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 2.50, y, 1.40, 0.32, sub, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 3.90, y, 1.55, 0.32, reg, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 5.45, y, 1.55, 0.32, ov, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 7.00, y, 2.70, 0.32, profile, size=9, italic=True, color=TXT_BODY,
             anchor=MSO_ANCHOR.MIDDLE)

# Bloc de lecture
add_rect(s, 0.30, 3.10, 9.40, 0.32, BG_CREAM)
add_text(s, 0.30, 3.10, 9.40, 0.32, "  Protocole : des dérives à date connue sur un bruit de fond réel",
         size=10, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.35, 3.50, 9.30, 1.65, [
    [("• Bruit de fond 100 % réel ; on injecte des dérives contrôlées (translations de moyenne, rotations de facteurs) à une date de rupture connue : le délai se mesure exactement.", 9, False, TXT_BODY)],
    [("• Fond neutralisé : on inverse au hasard le signe de journées entières de la coupe. Volatilité et corrélations réelles restent, toute dérive réelle disparaît. On obtient autant de trajectoires sans dérive que nécessaire pour compter les fausses alarmes.", 9, False, TXT_BODY)],
    [("• Mesure : délai moyen de détection (ADD) à fréquence de fausses alarmes fixée (ARL₀), comparé au plancher de Sharpe et au pooling.", 9, False, TXT_BODY)],
    [("• Validation réelle : phases haussières et baissières datées a posteriori (Pagan & Sossounov, 2003). BAF est synthétique, généré à partir d'un jeu réel anonymisé de fraude à l'ouverture de compte.", 9, False, TXT_BODY)],
])

# VISUEL À DESSINER À LA MAIN : série réelle neutralisée, avec une dérive injectée à la date de rupture τ* marquée par une ligne verticale.

footer_refs(s, "Jesus et al. (2022) · Pagan & Sossounov (2003) · Article B")

# =========================================================================
# SLIDE 9 — Conférences cibles, par échéance
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Conférences cibles, par échéance", 9)
badge(s, "ANALYSE", color=BLUE, w=1.00)

# Table header
add_rect(s, 0.30, 0.75, 9.40, 0.35, TEAL)
headers = [(0.30, 2.20, "Conférence"),
           (2.50, 1.40, "Rang"),
           (3.90, 1.55, "Échéance (est.)"),
           (5.45, 1.55, "Adéquation"),
           (7.00, 2.70, "Rôle")]
for x, w, text in headers:
    add_text(s, x, 0.75, w, 0.35, text, size=10, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Rows
data_rows = [
    ("FINEC 2027 (Lancaster)", "Finance*", "1er déc. 2026", "★★", "Premiers retours d'économètres", BG_CREAM, ORANGE),
    ("ICML 2027", "A*", "22 janv. 2027", "★★★", "Cible principale", BG_GREEN_LT, GREEN),
    ("SoFiE 2027", "Finance*", "≈ fin janv. 2027", "★★", "Visibilité en économétrie", BG_CREAM, ORANGE),
    ("UAI 2027", "A", "≈ 12 févr. 2027", "★★★", "Repli si janvier glisse", BG_GREEN_LT, GREEN),
    ("ICLR 2027 (ateliers)", "Workshop", "≈ févr. 2027", "★★", "Visibilité ML", BG_BLUE_LT, BLUE),
    ("NeurIPS 2027", "A*", "≈ mi-mai 2027", "★★★", "Repli après décision ICML", BG_GREEN_LT, GREEN),
]
for i, (conf, sub, reg, ov, profile, bg, color_text) in enumerate(data_rows):
    y = 1.10 + i * 0.32
    add_rect(s, 0.30, y, 9.40, 0.32, bg)
    add_text(s, 0.35, y, 2.20, 0.32, conf, size=9, bold=True, color=color_text,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 2.50, y, 1.40, 0.32, sub, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 3.90, y, 1.55, 0.32, reg, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 5.45, y, 1.55, 0.32, ov, size=9, color=TXT_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 7.00, y, 2.70, 0.32, profile, size=9, italic=True, color=TXT_BODY,
             anchor=MSO_ANCHOR.MIDDLE)

# Bloc de lecture
add_rect(s, 0.30, 3.10, 9.40, 0.32, BG_CREAM)
add_text(s, 0.30, 3.10, 9.40, 0.32, "  Lecture stratégique",
         size=10, bold=True, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 0.35, 3.50, 9.30, 1.65, [
    [("• Écartées : KDD 2027 (cycle 2) et ECML-PKDD 2027, voies possibles de l'Article B ; IJCAI 2027 (11 janv.), exclusive avec ICML ; ICAIF 2027 (≈ août), hors fenêtre.", 9, False, TXT_BODY)],
    [("• Filets : ECAI 2027 (14 avr., A), ICDM 2027 (≈ début juin, A*), CIKM et DSAA 2027 (≈ fin mai).", 9, False, TXT_BODY)],
    [("• * Conférences de finance sans actes : compatibles avec une relecture ICML en cours, règle à confirmer dans l'appel ICML 2027.", 9, False, TXT_BODY)],
    [("• ★ Si ICML refuse, les décisions tombent avant l'échéance NeurIPS : on resoumet avec les relectures.", 9, True, GREEN)],
])

footer_refs(s, "Rangs CORE · ★★★ viser · ★★ option · est. = estimée d'après l'édition 2026 · * sans actes")

# =========================================================================
# SLIDE 10 — Prochaines étapes
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Prochaines étapes", 10)
badge(s, "ACTIONS", color=ORANGE, w=1.00)

add_rect(s, 0.30, 0.75, 9.40, 0.35, BG_BLUE_LT)
add_text(s, 0.30, 0.75, 9.40, 0.35,
         "  D'ici deux semaines : un premier test qui peut réfuter l'idée",
         size=11, bold=True, color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)

points = [
    (TEAL, "Écrire la théorie",
     "Borne de délai dans l'espace latent et calibration exacte des moniteurs latents ; chaque borne vérifiée par simulation."),
    (BLUE, "Auditer le panel S&P 500",
     "Part de variance du marché, corrélations, et information de dérive des régimes passés, direction par direction."),
    (ORANGE_LT, "Construire moniteurs et références",
     "Surveillance univariée, pooling et PCA-orthogonal, branchés sur la même pile de détection."),
    (RED, "Premier test d'échappement",
     "Sur bruit de fond réel neutralisé, le gain doit continuer de croître avec le nombre d'actifs là où le pooling plafonne. Sinon, l'idée est à revoir."),
]
for i, (color, title, desc) in enumerate(points):
    y = 1.30 + i * 0.85
    vertical_bar(s, 0.40, y, color=color, height=0.65)
    add_text(s, 0.65, y - 0.02, 9.00, 0.28, title, size=10, bold=True, color=TXT_DARK)
    add_text(s, 0.65, y + 0.26, 9.00, 0.45, desc, size=9, color=TXT_BODY)

footer_refs(s, "Article B (The Whitening Advantage) · Pagan & Sossounov (2003)")

# =========================================================================
# SLIDE 11 — DIVIDER — Questions et discussion
# =========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0.00, 1.50, 10.00, 2.60, TEAL)
add_rect(s, 1.20, 1.85, 1.50, 1.50, RED)
add_text(s, 1.20, 1.85, 1.50, 1.50, "?", size=42, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_paragraphs(s, 3.20, 1.90, 6.00, 1.00, [
    [("Questions &", 24, True, WHITE)],
    [("discussion", 24, True, WHITE)],
])
add_text(s, 3.20, 3.00, 6.00, 0.80,
         "Contact : R. Minato · Code : dépôt public à venir, anonymisé pendant la relecture",
         size=12, color=RGBColor(0xD5, 0xDB, 0xDB))
page_num_only(s, 11)

# =========================================================================
# SLIDE 12 — ANNEXE — D'où vient le plancher de Sharpe
# =========================================================================
s = prs.slides.add_slide(BLANK)
header_bar(s, "Annexe : d'où vient le plancher de Sharpe", 12)
badge(s, "ANNEXE", color=TEAL, w=0.95)

add_rect(s, 0.30, 0.75, 9.40, 0.30, BG_BLUE_LT)
add_text(s, 0.30, 0.75, 9.40, 0.30, "  Quatre lignes, pour une dérive gaussienne de la moyenne",
         size=10, bold=True, color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)

# Étapes numérotées (ligne en gras + texte simplifié en italique)
add_paragraphs(s, 0.35, 1.15, 9.30, 0.50,[
    [("1. Preuve quotidienne d'un décalage μ de la moyenne, à volatilité σ :", 10, True, TXT_DARK)],
    [("KL = μ² / (2σ²) = SR_jour² / 2, avec SR_jour = μ/σ", 9, False, TXT_DARKER, True)],
])

add_paragraphs(s, 0.35, 1.85, 9.30, 0.50,[
    [("2. Annualisation sur 252 jours de bourse :", 10, True, TXT_DARK)],
    [("SR = √252 · SR_jour ⇒ KL = SR² / 504", 9, False, TXT_DARKER, True)],
])

add_paragraphs(s, 0.35, 2.60, 9.30, 0.50, [
    [("3. Borne de la détection la plus rapide, atteinte par le CUSUM ; (1 + o(1)) s'efface quand ARL₀ grandit :", 10, True, TXT_DARK)],
    [("ADD ≥ ln(ARL₀) / KL · (1 + o(1))", 9, False, TXT_DARKER, True)],
])

add_paragraphs(s, 0.35, 3.30, 9.30, 0.50, [
    [("4. D'où le plancher, avec α₀ = 1/ARL₀ :", 10, True, TXT_DARK)],
    [("ADD_min ≥ 504 · ln(1/α₀) / SR²  □", 9, False, TXT_DARKER, True)],
])

# Bloc rose — portée
add_rect(s, 0.30, 4.05, 9.40, 1.05, BG_PINK)
add_paragraphs(s, 0.40, 4.10, 9.20, 0.95, [
    [("Portée :", 10, True, RED)],
    [("• La borne vise la loi inconditionnelle d'une seule série. Un moniteur qui connaîtrait la volatilité du jour ferait mieux : l'Article B le mesure avec un oracle à anticipation.", 9, False, TXT_BODY)],
    [("• Le moniteur de signe de l'Article B paie un facteur π/2 sous loi gaussienne, en échange d'une calibration exacte.", 9, False, TXT_BODY)],
    [("• En multivarié, SR² devient le carré du Sharpe du portefeuille aligné sur la dérive : sujet de la semaine prochaine.", 9, False, TXT_BODY)],
])
footer_refs(s, "Lorden (1971) · Moustakides (1986) · Lai (1998) · Tartakovsky et al. (2014)")

# =========================================================================
# SAVE
# =========================================================================
out = "/home/m53/2026_10_12_Point_Hebdo_JEPA_Drift.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Total slides: {len(prs.slides)}")
